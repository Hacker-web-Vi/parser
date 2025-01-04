import json
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
    
json_file_path = "metrics-end-day.json"
output_xlsx_path = "all-metrics.xlsx"

with open(json_file_path, "r") as file:
    metrics = json.load(file)

workbook: Workbook = openpyxl.Workbook()

if "Sheet" in workbook.sheetnames:
    workbook.remove(workbook["Sheet"])

sheet: Worksheet = workbook.create_sheet(title="Overall")

center_alignment = Alignment(horizontal="center", vertical="center")

bold_border = Border(
    left=Side(border_style="medium", color="000000"),
    right=Side(border_style="medium", color="000000"),
    top=Side(border_style="medium", color="000000"),
    bottom=Side(border_style="medium", color="000000")
)

header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
###

headers = ["Moniker", "Validator Address", "Tombstoned", "Slashes #", "Active #", "Proposed #", "Signed #", "Missed #", "Uptime", "Avg. Uptime per day"]

header_font = Font(bold=True, size=12)
range_font = Font(bold=False, size=10, color='707070')

for col_num, header in enumerate(headers, start=1):
    header_cell = sheet.cell(row=1, column=col_num, value=header)
    
    header_cell.font = header_font
    header_cell.alignment = center_alignment
    header_cell.fill = header_fill
    header_cell.border = bold_border

sorted_days = sorted(metrics['day_boundaries'].items(), key=lambda x: x[1]['start'])

day_to_col_num = {}
col_num = len(headers) + 1

for day, day_data in sorted_days:
    day_cell = sheet.cell(row=1, column=col_num, value=day)
    day_cell.font = header_font
    day_cell.alignment = center_alignment
    day_cell.fill = header_fill
    day_cell.border = bold_border

    day_boundaries = f"{day_data['start']} - {day_data['end']}"
    day_boundaries_cell = sheet.cell(row=2, column=col_num, value=day_boundaries)
    day_boundaries_cell.font = range_font
    day_boundaries_cell.alignment = center_alignment
    day_boundaries_cell.border = bold_border

    sheet.merge_cells(start_row=1, start_column=col_num, end_row=1, end_column=col_num + 1)
    sheet.merge_cells(start_row=2, start_column=col_num, end_row=2, end_column=col_num + 1)

    day_to_col_num[day] = col_num
    col_num += 2
    
row_num = 3
sorted_validators = sorted(metrics["validators"], key=lambda x: x["total_signed_blocks"] + x["total_missed_blocks"], reverse=True)

for validator in sorted_validators:
    moniker = validator["moniker"]
    valoper = validator["valoper"]
    tombstoned = str(validator["tombstoned"])
    slashes = len(validator['slashes'])
    total_proposed = validator["total_proposed_blocks"]
    total_signed = validator["total_signed_blocks"]
    total_missed = validator["total_missed_blocks"]
    total_uptime = round((
        (total_signed / (total_signed + total_missed))
        if (total_signed + total_missed) > 0 else 0.0
    ),5)

    total_active = total_signed + total_missed
    
    uptime_days = []
    for day in validator['dates']:
        signed = validator['dates'][day]['signed_count']
        missed = validator['dates'][day]['missed_count']
        day_uptime = round((
            (signed / (signed + missed))
            if (signed + missed) > 0 else 0.0
        ),5)
        
        if (signed + missed) > 0:
            uptime_days.append(day_uptime)

            light_blue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
            light_red_fill = PatternFill(start_color="FFCCCB", end_color="FFCCCB", fill_type="solid")
            light_green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

            balck_bold_fond = Font(color="000000", bold=True)

            signed_cell = sheet.cell(row=row_num, column=day_to_col_num[day], value=signed)
            signed_cell.font = balck_bold_fond
            signed_cell.alignment = center_alignment
            signed_cell.fill = light_green_fill
            
            missed_cell = sheet.cell(row=row_num, column=day_to_col_num[day] + 1, value=missed)
            missed_cell.font = balck_bold_fond
            missed_cell.alignment = center_alignment
            missed_cell.fill = light_red_fill


            uptime_cell = sheet.cell(row=row_num+1, column=day_to_col_num[day], value=day_uptime)
            uptime_cell.font = balck_bold_fond
            uptime_cell.alignment = center_alignment
            uptime_cell.fill = light_blue_fill
            uptime_cell.border = bold_border
            uptime_cell.number_format = '0.000%'


            sheet.merge_cells(start_row=row_num+1, start_column=day_to_col_num[day], end_row=row_num+1, end_column=day_to_col_num[day] + 1)
    
    days_uptime = round(sum(uptime_days) / len(uptime_days), 5) if uptime_days else 0.00000

    cell_moniker = sheet.cell(row=row_num, column=1, value=moniker)
    cell_moniker.alignment = center_alignment
    cell_moniker.border = bold_border

    cell_valoper = sheet.cell(row=row_num, column=2, value=valoper)
    cell_valoper.alignment = center_alignment
    cell_valoper.border = bold_border

    cell_tombstoned = sheet.cell(row=row_num, column=3, value=tombstoned)
    false_font = Font(bold=True, size=12, color='008000')
    true_font = Font(bold=True, size=12, color='FF0000')
    cell_tombstoned.font = true_font if tombstoned == 'True' else false_font
    cell_tombstoned.alignment = center_alignment
    cell_tombstoned.border = bold_border

    cell_slashes = sheet.cell(row=row_num, column=4, value=slashes)
    cell_slashes.font = true_font if slashes else false_font
    cell_slashes.alignment = center_alignment
    cell_slashes.border = bold_border

    cell_total_active = sheet.cell(row=row_num, column=5, value=total_active)
    cell_total_active.alignment = center_alignment
    cell_total_active.border = bold_border

    cell_total_proposed = sheet.cell(row=row_num, column=6, value=total_proposed)
    cell_total_proposed.alignment = center_alignment
    cell_total_proposed.border = bold_border

    cell_total_signed = sheet.cell(row=row_num, column=7, value=total_signed)
    cell_total_signed.alignment = center_alignment
    cell_total_signed.border = bold_border

    cell_total_missed = sheet.cell(row=row_num, column=8, value=total_missed)
    cell_total_missed.alignment = center_alignment
    cell_total_missed.border = bold_border

    cell_total_uptime = sheet.cell(row=row_num, column=9, value=total_uptime)
    cell_total_uptime.alignment = center_alignment
    cell_total_uptime.border = bold_border
    cell_total_uptime.number_format = '0.000%'

    cell_days_uptime = sheet.cell(row=row_num, column=10, value=days_uptime)
    cell_days_uptime.alignment = center_alignment
    cell_days_uptime.border = bold_border
    cell_days_uptime.number_format = '0.000%'

    sheet.merge_cells(start_row=row_num, start_column=1, end_row=row_num + 1, end_column=1)
    sheet.merge_cells(start_row=row_num, start_column=2, end_row=row_num + 1, end_column=2)
    sheet.merge_cells(start_row=row_num, start_column=3, end_row=row_num + 1, end_column=3)
    sheet.merge_cells(start_row=row_num, start_column=4, end_row=row_num + 1, end_column=4)
    sheet.merge_cells(start_row=row_num, start_column=5, end_row=row_num + 1, end_column=5)
    sheet.merge_cells(start_row=row_num, start_column=6, end_row=row_num + 1, end_column=6)
    sheet.merge_cells(start_row=row_num, start_column=7, end_row=row_num + 1, end_column=7)
    sheet.merge_cells(start_row=row_num, start_column=8, end_row=row_num + 1, end_column=8)
    sheet.merge_cells(start_row=row_num, start_column=9, end_row=row_num + 1, end_column=9)
    sheet.merge_cells(start_row=row_num, start_column=10, end_row=row_num + 1, end_column=10)

    row_num += 2

workbook.save(output_xlsx_path)

print(f"Excel file created: {output_xlsx_path}")
