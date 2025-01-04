import json
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from datetime import datetime

# def main(metrics, days, output_xlsx_path, sheet_name):
#     workbook: Workbook = openpyxl.Workbook()

#     if "Sheet" in workbook.sheetnames:
#         workbook.remove(workbook["Sheet"])

#     sheet: Worksheet = workbook.create_sheet(title=sheet_name)

#     center_alignment = Alignment(horizontal="center", vertical="center")

def main(metrics, days, output_xlsx_path, sheet_name, valopers):
    try:
        workbook = openpyxl.load_workbook(output_xlsx_path)
    except FileNotFoundError:
        workbook = openpyxl.Workbook()

    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])

    if sheet_name in workbook.sheetnames:
        workbook.remove(workbook[sheet_name])

    sheet: Worksheet = workbook.create_sheet(title=sheet_name)

    center_alignment = Alignment(horizontal="center", vertical="center")

    bold_border = Border(
        left=Side(border_style="medium", color="000000"),
        right=Side(border_style="medium", color="000000"),
        top=Side(border_style="medium", color="000000"),
        bottom=Side(border_style="medium", color="000000")
    )

    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    headers = ["Moniker", "Validator Address", "Tombstoned", "Slashes #", "Active #", "Proposed #", "Signed #", "Missed #", "Uptime", "Avg. Uptime per day"]

    header_font = Font(bold=True, size=12)
    range_font = Font(bold=False, size=10, color='707070')

    for col_num, header in enumerate(headers, start=1):
        header_cell = sheet.cell(row=1, column=col_num, value=header)
        
        header_cell.font = header_font
        header_cell.alignment = center_alignment
        header_cell.fill = header_fill
        header_cell.border = bold_border

    sorted_days = sorted(metrics['day_boundaries'].items(), key=lambda x: x[1]['start'])

    day_to_col_num = {}
    col_num = len(headers) + 1

    for day, day_data in sorted_days:
        if day in days:
            day_cell = sheet.cell(row=1, column=col_num, value=day)
            day_cell.font = header_font
            day_cell.alignment = center_alignment
            day_cell.fill = header_fill
            day_cell.border = bold_border

            day_boundaries = f"{day_data['start']} - {day_data['end']}"
            day_boundaries_cell = sheet.cell(row=2, column=col_num, value=day_boundaries)
            day_boundaries_cell.font = range_font
            day_boundaries_cell.alignment = center_alignment
            day_boundaries_cell.border = bold_border

            sheet.merge_cells(start_row=1, start_column=col_num, end_row=1, end_column=col_num + 1)
            sheet.merge_cells(start_row=2, start_column=col_num, end_row=2, end_column=col_num + 1)

            day_to_col_num[day] = col_num
            col_num += 2
        
    row_num = 3
    sorted_validators = sorted(metrics["validators"], key=lambda x: x["total_signed_blocks"] + x["total_missed_blocks"], reverse=True)
    
    for validator in sorted_validators:
        if validator['valoper'] not in valopers:
            continue
        
        moniker = validator["moniker"]
        valoper = validator["valoper"]
        tombstoned = str(validator["tombstoned"])
        slashes_list = [slash for slash in validator['slashes'] if slash['time'].split('T')[0] in days]
        slashes = len(slashes_list)
        total_proposed  = 0
        total_signed  = 0
        total_missed  = 0
        
        uptime_days = []
        for day in validator['dates']:
            if day in days:
                signed = validator['dates'][day]['signed_count']
                missed = validator['dates'][day]['missed_count']
                proposed = validator['dates'][day]['proposed_count']
                total_proposed += proposed
                total_missed += missed
                total_signed += signed

                day_uptime = round((
                    (signed / (signed + missed))
                    if (signed + missed) > 0 else 0.0
                ),5)
                
                if (signed + missed) > 0:
                    uptime_days.append(day_uptime)

                    light_blue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
                    light_red_fill = PatternFill(start_color="FFCCCB", end_color="FFCCCB", fill_type="solid")
                    light_green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

                    balck_bold_fond = Font(color="000000", bold=True)

                    signed_cell = sheet.cell(row=row_num, column=day_to_col_num[day], value=signed)
                    signed_cell.font = balck_bold_fond
                    signed_cell.alignment = center_alignment
                    signed_cell.fill = light_green_fill
                    
                    missed_cell = sheet.cell(row=row_num, column=day_to_col_num[day] + 1, value=missed)
                    missed_cell.font = balck_bold_fond
                    missed_cell.alignment = center_alignment
                    missed_cell.fill = light_red_fill


                    uptime_cell = sheet.cell(row=row_num+1, column=day_to_col_num[day], value=day_uptime)
                    uptime_cell.font = balck_bold_fond
                    uptime_cell.alignment = center_alignment
                    uptime_cell.fill = light_blue_fill
                    uptime_cell.border = bold_border
                    uptime_cell.number_format = '0.000%'


                    sheet.merge_cells(start_row=row_num+1, start_column=day_to_col_num[day], end_row=row_num+1, end_column=day_to_col_num[day] + 1)

        total_uptime = round((
            (total_signed / (total_signed + total_missed))
            if (total_signed + total_missed) > 0 else 0.0
        ),5)

        
        total_active = total_signed + total_missed

        days_uptime = round(sum(uptime_days) / len(uptime_days), 5) if uptime_days else 0.00000

        cell_moniker = sheet.cell(row=row_num, column=1, value=moniker)
        cell_moniker.alignment = center_alignment
        cell_moniker.border = bold_border

        cell_valoper = sheet.cell(row=row_num, column=2, value=valoper)
        cell_valoper.alignment = center_alignment
        cell_valoper.border = bold_border

        cell_tombstoned = sheet.cell(row=row_num, column=3, value=tombstoned)
        false_font = Font(bold=True, size=12, color='008000')
        true_font = Font(bold=True, size=12, color='FF0000')
        cell_tombstoned.font = true_font if tombstoned == 'True' else false_font
        cell_tombstoned.alignment = center_alignment
        cell_tombstoned.border = bold_border

        cell_slashes = sheet.cell(row=row_num, column=4, value=slashes)
        cell_slashes.font = true_font if slashes else false_font
        cell_slashes.alignment = center_alignment
        cell_slashes.border = bold_border

        cell_total_active = sheet.cell(row=row_num, column=5, value=total_active)
        cell_total_active.alignment = center_alignment
        cell_total_active.border = bold_border

        cell_total_proposed = sheet.cell(row=row_num, column=6, value=total_proposed)
        cell_total_proposed.alignment = center_alignment
        cell_total_proposed.border = bold_border

        cell_total_signed = sheet.cell(row=row_num, column=7, value=total_signed)
        cell_total_signed.alignment = center_alignment
        cell_total_signed.border = bold_border

        cell_total_missed = sheet.cell(row=row_num, column=8, value=total_missed)
        cell_total_missed.alignment = center_alignment
        cell_total_missed.border = bold_border

        cell_total_uptime = sheet.cell(row=row_num, column=9, value=total_uptime)
        cell_total_uptime.alignment = center_alignment
        cell_total_uptime.border = bold_border
        cell_total_uptime.number_format = '0.000%'

        cell_days_uptime = sheet.cell(row=row_num, column=10, value=days_uptime)
        cell_days_uptime.alignment = center_alignment
        cell_days_uptime.border = bold_border
        cell_days_uptime.number_format = '0.000%'

        sheet.merge_cells(start_row=row_num, start_column=1, end_row=row_num + 1, end_column=1)
        sheet.merge_cells(start_row=row_num, start_column=2, end_row=row_num + 1, end_column=2)
        sheet.merge_cells(start_row=row_num, start_column=3, end_row=row_num + 1, end_column=3)
        sheet.merge_cells(start_row=row_num, start_column=4, end_row=row_num + 1, end_column=4)
        sheet.merge_cells(start_row=row_num, start_column=5, end_row=row_num + 1, end_column=5)
        sheet.merge_cells(start_row=row_num, start_column=6, end_row=row_num + 1, end_column=6)
        sheet.merge_cells(start_row=row_num, start_column=7, end_row=row_num + 1, end_column=7)
        sheet.merge_cells(start_row=row_num, start_column=8, end_row=row_num + 1, end_column=8)
        sheet.merge_cells(start_row=row_num, start_column=9, end_row=row_num + 1, end_column=9)
        sheet.merge_cells(start_row=row_num, start_column=10, end_row=row_num + 1, end_column=10)

        row_num += 2

    workbook.save(output_xlsx_path)
    print(f"Excel file created: {output_xlsx_path}")

if __name__ == '__main__':

    ROUND_1_VALOPERS = ['storyvaloper12s7mczq0gsvzyav5sleje29pped8mpgcprr9r2', 'storyvaloper1qzlsuhxyggc58dakmk92420slnsares5c3mk6k', 'storyvaloper1c58gr53xrf6ktwtr8vwmzxyfz2xzzaaz8rrw68', 'storyvaloper1hvavwqtnjw5m3vtkzgm62kwsfjkhwq7h8m4azm', 'storyvaloper1l6gpgqmucdn7yyuvdyeph74cgs2f7k66u60hq6', 'storyvaloper13el67lyw2qpacpuzha9e7p3rpjg8jl8tmecq2q', 'storyvaloper1zrkjleyv9c4rp5u24cyn6j08ahypep35r4uk3k', 'storyvaloper12s2acumq6ycr80jycpt2hzq94xqgtehv2dz8h2', 'storyvaloper1htkcu0adnlfqg4l29t2n5cc6l2nywle6qyqdr9', 'storyvaloper13tm6a04err3yzasegk4agqgqfgtt048yhs8hs6', 'storyvaloper1ml7wkkeuupt9syle0zszvst88632mz8g3aahpy', 'storyvaloper1mg5wjz3kta5dfqxc93hv9hjg0cha53l97ne3z4', 'storyvaloper1the6ag95nrgg0h6pw3nzxdf4zv4e4uws0x3yln', 'storyvaloper1hcndmwxa0dmdvkzxlxlvdxd033n5jcmh6afkjh', 'storyvaloper13k24q8a30twnuzkpsk9c95xhs7cwdfckx9d9e7', 'storyvaloper1tnhycpfuzsm4l7l6t3qsy5j2lged8ec8eyft2k', 'storyvaloper1rfnera4cjh83fahr62dlfz2n9yf42v45nltkwy', 'storyvaloper1wl873hxdkhhjalt86zgmyphfzlt2kl39kf4x2a', 'storyvaloper12vrh9n58r2ftpdjfjj7j4ghfw76sx9c4e45yxl', 'storyvaloper14vnfjngu4y2ywga0aunvtaqupngqt2cr9qwz22', 'storyvaloper1u30x96px5t3q3kl4n6d0d5sl52zkw5vqylp9fz', 'storyvaloper1xx42vsmgccj96hnfw5tumtf062c09z678pg46l', 'storyvaloper1ymsmqlx54g33msau9x98my03p9hd36djvkylqd', 'storyvaloper1zsw5cflczjxrk5ps5lj20dzua94pjdn72guhyn', 'storyvaloper19320qvhvyy0gqgmxjewt6hrpmun5tac6sp2f76', 'storyvaloper1xa3rd3vs28jym7pwcacf2t3cgpl590zj6eyuhq', 'storyvaloper1qxhdqgsmm6vlrfrn594mm46qqenu52cckwa6cc', 'storyvaloper1l7wq80kk7j7e84ml7mzdhrxd66xufm7pvunwj6', 'storyvaloper1z0uqv8xrw2j23mlcpmzatt3jj9hzezdq08ar8t', 'storyvaloper1cu3ua6279lk9w2kxlulzr2ttyzpxd5hytyfc4y', 'storyvaloper1pdj0z84lau0l7vf2jl4qs7yggv48p3avy82spr', 'storyvaloper1kz2cgurfhgege6zqnk8vwupwzqcjry0du23g4z', 'storyvaloper1d0k6pzc6kpvfwguqwxer87aef5njrqngdgasu0', 'storyvaloper19t6g3988z2zhwv9xzzdn9zn0ddsserwln3sjn8', 'storyvaloper1765xe28ktlpkxvwsndx4aswx2rwxt24v7lx4u7', 'storyvaloper16xg48zr4m6unpwqd6c9q2uj4t3nrllqp8tw02u', 'storyvaloper1x4e6l5dgfap6yhpdeg4a28dad0lfgg39nypxlt', 'storyvaloper1dzw4c607cnqnxe0djvr5a8chjm0h99yjxl7tpy', 'storyvaloper1d055v5vthzzf37n5rltkr2luka7xulq95w25a4', 'storyvaloper13cm9skv9yq6s983axr4hyy45awcaxtwtqsj8tu', 'storyvaloper1wte6l2r6xh3l6xus8ug0u36xya0cltclyuqy75', 'storyvaloper1s5j9d80stg70pdg2jan9ea5ln9zcdraarqjark', 'storyvaloper187xp6lza085qvj28cpnp55vaexldtlnggkplkx', 'storyvaloper1upyt7dr96p2ffl62ecwypeglqvk7596a3xrmvn', 'storyvaloper10llgt50kejk7exvsvkr5wysr7vqdrmdmwfd37e', 'storyvaloper1cvsdp0tsz25fhedd7cjvntq42347astvar06v8', 'storyvaloper1enja57p5uhknfynelkkvc64nraeh3hy4p9rxz9', 'storyvaloper16exhdc2hzrtralnk5meca08nd8kfeu00kvxgs3', 'storyvaloper15a4fph700vr0t0arfsrleeyh4al9087ee3evct', 'storyvaloper16nwswhexfkv36f57um68td5j7mtsf9j4wfj4dn', 'storyvaloper1l6m9y6l064mp7hvzt3vqg3fnuu6vx0lkn50udf', 'storyvaloper18m7p3puy2scyqp83369nhlmmerk3vmgs4rngfc', 'storyvaloper1v8pyhyc8npkkl04tpt7uwmev3plpndvwh04qq7', 'storyvaloper1r62nlglv96dqhqhsc2dq3tfydh7kp2pq8f7w07', 'storyvaloper1f55pj65qp0rufg2y5u47hz98gtf6jzpts8lpgy', 'storyvaloper1wh896802py6w66ymexxqdduf0d4jq4zw04yqzr', 'storyvaloper1g8lyk7xtfxahjdkta5jcazwc2pz7ht8ucj9wzl', 'storyvaloper12e6dh28hq7hd2qwd6v992uwr5ukrmq2xeup4ep', 'storyvaloper13hzy82h0cg9xjc0xhc8u28de92ntdjlae3te6u', 'storyvaloper1999gqujrmash4r8k256vdgpurd3h7067dvxtun', 'storyvaloper1x9c7xr8x4du2e926cgztthaq8cydnvcvvypesa', 'storyvaloper1ecsq9zy0huqvjkuvqewtzea3dtvc5j4hzqju9a', 'storyvaloper1d0g7xp0qu2dxlqps65cfzz4hx98tncs57qa95t', 'storyvaloper1qxs6uh58zlh0pc59er06xvj53rz0mn9hpfgv25', 'storyvaloper1cn4agxuv5zl96tw3l89w9x89rv6mjcqadcx9y3', 'storyvaloper1spgujf9pjjtp8qrlvuucdd6cxfe7xs7858k7z5', 'storyvaloper1ttetgy8a5lcr78py460pm6cjr83jcnud26jff9', 'storyvaloper1ckj7xq5uk44tqh9sty8qct9cz3hemh0zgyktar', 'storyvaloper1sp0u3967vczunl25sxh8ftk7hrex8qe74czqks', 'storyvaloper1wjl6dzufze67r7rpx5ta2z7kxeeva4h2sawxa8', 'storyvaloper16qcd4vf5y8getfsc2a8amww8kawqrcwpmwrpua', 'storyvaloper1fh06ygln2lndkvnvswrtc9v2j82m58tzceux6u', 'storyvaloper1kykx4cc73cwrnqus2rflullfw67mrwdde73xp5', 'storyvaloper15m5s09ptvuyrmmml5g7e2xg7hrewp3x5cdufu4', 'storyvaloper1xcmr77ya82eqjp2nkcv5s9mdcuh2tmxav9d9l7', 'storyvaloper1s5gd8td90j2cgn9h9we2zpnn0kesespp7nfpfv', 'storyvaloper1c74pd36ddkr7kk3jtrxw9qhck6v4gn2cdtg3m4', 'storyvaloper13kvhv7fny79pdnksuq229ndc0ylnl3wcc697da', 'storyvaloper1achxz26kc2mn8458vc2evllnrd9s2mr42vyajk', 'storyvaloper1l3t54lspt8pwxzfncg8t5flx4yarkjlj0tpc53', 'storyvaloper1egge9775vhypgsgd9mqq2uw3q5gqut708hj57y', 'storyvaloper1pjfazvhc93m5s7jyx4md36nxxllmhedkt77wc7', 'storyvaloper14n0ksrhhk2zdrypculd4zuaa64j4cueczq5xhm', 'storyvaloper1fuxw9xujuw08cw6ujqx6upw6ja8zcerwsdn224', 'storyvaloper18j8r0grm4yh7lfqvftaxdurum2gdnh567yt3s0', 'storyvaloper18cvm8frrckjwfshqdg6p63pzwy9dxvsrlup8kk', 'storyvaloper1gnkssjh2nva4z0yvv7megl25hmmvg9csvggt0h', 'storyvaloper12smnjtdaxdks6zhkxddg3jw5fj9lsruzauwjap', 'storyvaloper1v7v69m6hdz3ydewanqwypzzxekrsqujqydu9n9', 'storyvaloper1826z0gntznjlcpa4weu55uvtp5mmd62ns05gyq', 'storyvaloper10p94srhdxunq2vcnx6t42zwl8365v0ac9jrvl7', 'storyvaloper16h45lp8ydnt5e5307n8r3dkzxnvunm548qxpgu', 'storyvaloper1wv2wuyk00y6fscxnyc5m7j9fr4vpwj4thq3kt9', 'storyvaloper1kwf9x7n8pmpvv7zrhxjfqqecz65a8qgunxaz0q', 'storyvaloper1u2vlklgdfjnah2ll7gghp2cvw4j7x7asynxxdn', 'storyvaloper16a8rgcx9287q0nusu7frw7n3yrk9hlpn0re0ht', 'storyvaloper1rrwuf534h4ft8r5pr7an38u89s3f7gjj839rrh', 'storyvaloper1n9ft9mx5q0amlwextftwpehxe082sj2sj8tzrk', 'storyvaloper1tcrwtm792svad9nfkxumyesakx80t2ha4txfns', 'storyvaloper1fsud00esx6rtp7n7p7mj2xg4v6l7w8lhy8ztdk', 'storyvaloper1mrgqjazkfmdj3e2tj4facwkshlzqym3h4g9tan', 'storyvaloper1xjvw3wj7kxvp2f4j5a0fud96rjy09dpxre6yq7', 'storyvaloper12vs37z0ee7m9xns8zkqsng9qnf5pkvcnqgdzqt', 'storyvaloper1f0sljxr2w2vexhmdytgjz6dw57vraadd973t00', 'storyvaloper1acv5yknlf7atwf03mxd2lecwelmgtctc8j6vcl', 'storyvaloper1u0jw7jnesvqr55fxfmjycm5uxv9w6mjtfzkwdf', 'storyvaloper1gm8mm8zkcxaz6nnue7j3dhh3fxxjh5nzmfnena', 'storyvaloper1v4me33grq0uys8f2s0k0dllshlljpzk5yz348y', 'storyvaloper1lgexam6a80s5yj270tdw7erc54s9xzat7yaama', 'storyvaloper13nwrgmldnp9g9s88cx74a32as7cle36u0v8ck7', 'storyvaloper1nsnf4gmwc46clml3w37nz2yd9340rj23qd5emc', 'storyvaloper1u0mqgma7h300dj72yyy87u8vrzev7h4dymxfyg']

    # ROUND_2_VALOPERS = [
    #     "storyvaloper1v06kvxy79w02cna4d93c387wzvuewwxnqyxdld",
    #     "storyvaloper14s26qxtpyqkwr76atnn8324tvrhs89j34l8ugf",
    #     "storyvaloper1xsrjrajnrpx7u2gn7msxz85z469505fvunvsn5",
    #     "storyvaloper1xthp0asrd4eypxcxc5pxc70ctqpfma88ek3vdj",
    #     "storyvaloper1udvjz9ayx65e9xh55fr78v8zn30jsmlrffwgd3",
    #     "storyvaloper1a2khawzy374cm42tfn6v7uglplxkkld6hzpkpv",
    #     "storyvaloper1daxk0lukad4950ga6t2nc939gjvpwranre5dd4",
    #     "storyvaloper139kxvvqfzejt2ejk9ptt9t96h78rxd73xrj927",
    #     "storyvaloper1p3gu02nw270k4v8uuvhkxdzheqw7gxaalhyzq6",
    #     "storyvaloper1tg3tzx9nt6069sxrw4s6e0g50h4ljmfnc0te4l",
    #     "storyvaloper1zwrnppxgpzm437se3j0asq5k2yk38hc758pdr7",
    #     "storyvaloper1v8pyhyc8npkkl04tpt7uwmev3plpndvwh04qq7",
    #     "storyvaloper10aax54nlscnz4ly62klcmmfs4ys3hgtp0lftag",
    #     "storyvaloper1kuxt95cnl7y8pwzw9us0tpe74e2shawsyj5a0x",
    #     "storyvaloper1faykyn72emusnskkhs2zflkfj80xnx36lek3q9",
    #     "storyvaloper1vq80dadh9zplaus3pkhasvqc3guztnxhj0fj38",
    #     "storyvaloper1ulrqpkgw5utu35362vpvknql9x6qresgags8q2",
    #     "storyvaloper14j90lm6qu2uaux4fkhvg9ftxnl6vff5cegz95r",
    #     "storyvaloper10ryxkfzlxz9l2e85x37k7lr24vd7hxfwz50w3s",
    #     "storyvaloper13tmg5f5mtdkj87vq94nlsmhu3kylg6askjvw2f",
    #     "storyvaloper1r6967ky90cvtpmzag5l2hyqqln7kmcetg72sne",
    #     "storyvaloper1vvzqyum2njq76w6d2cqeh5a4upqz44g25nzc0s",
    #     "storyvaloper1ujeanmqgxjth2cmm4mmf906pcmvz6urk0htsfv",
    #     "storyvaloper1mu457ah3ry685w3f0pnq940mpdp04j0skqwtt9",
    #     "storyvaloper1yj8ymdm7l2ek9cjevs3kqxv06gysk6vh2564m3",
    #     "storyvaloper17y27wj0n54uajlsrt7umquq56dfe966ljfzx7g",
    #     "storyvaloper1zqxgv84754uy2nz6axe50vd2us0xgh7gghzrz2",
    #     "storyvaloper1rzhgd5rdcff0uee3jahp75guypgalkuldm2tq2",
    #     "storyvaloper1rxa60lwchrr0mzsrwqud6rrz6jgdat7gsglmet",
    #     "storyvaloper1dnhkp552qr7nwr30nu563uzau7qv5dl9myzt4r",
    #     "storyvaloper14vnfjngu4y2ywga0aunvtaqupngqt2cr9qwz22",
    #     "storyvaloper1qj87mt3zwkfs088430dghtgp8n7cdjs49jvem2",
    #     "storyvaloper17s93lgl2ym8f98dukr7n9524sedxh8kjgkyfmv",
    #     "storyvaloper137cw3ke8e7lq8u0r57yxqatc6w2m9k379k0zpa",
    #     "storyvaloper1rquq9mmuf9wdy04jvlp5ur0f9nkkwxynft3h7y",
    #     "storyvaloper17u46ud0a7wty94yyltqpqkr3fdq9ft9rw0wsgs",
    #     "storyvaloper1fhrm40tnvgf4qp5ywn74pwn72t5eynrgdljds3",
    #     "storyvaloper16nwswhexfkv36f57um68td5j7mtsf9j4wfj4dn",
    #     "storyvaloper1ymjw6ydnv35t4dn8v8p7sz092z9fgxj8sltyqd",
    #     "storyvaloper1d4kf8mu5e84xkqespxc87z5ds8yf26t7p3apqd",
    #     "storyvaloper12wt5zqursxdkkvjhp4vrg5qkrjxyn0jaltn6yh",
    #     "storyvaloper1yn4gzfjpzuwjpva8e5tspjwufz3mjc0xe07mv5",
    #     "storyvaloper16e9gt2ppsfysu0sujh8320yffkgww6h4xg2vkq",
    #     "storyvaloper1qlglvc39s9xmu06x2tqulhvfh56xj93503ep6x",
    #     "storyvaloper1amtl58gyd69sy6kdlex9kn5dfc5vkntds539ya",
    #     "storyvaloper1an0thyxr7sgj4ar629whcaeejwkd3duma5djmr",
    #     "storyvaloper13jhfwsvhytq6qw3e42yalg7h88sgtecq46dgu2",
    #     "storyvaloper18s6ahmdc27yl3390ptjt98qr03h8tdjnv9hakz",
    #     "storyvaloper1r9642cstr2z4rwt4z2fgvk50zfcu7gt3r87kgg",
    #     "storyvaloper1rtagk36fkxsa9ntdt6fzq5v2cnnmjhw45zrq96",
    #     "storyvaloper1gdmu99t8qe29jp0a25k0dr9rpmvgxjt0dhc6xl",
    #     "storyvaloper1auna7t7rdcywwq7dlcxy6wn0h8y7y4gl8mxm0k",
    #     "storyvaloper1fqss80elgqdgjz5gtw2dyl82ayh0dy9zz58s47",
    #     "storyvaloper17r05q6nvmzdqnhwd2hv5zguzwulhgfp9a0hd3w",
    #     "storyvaloper1e7ddt2nsgagglccj3tgj8gk4n6rye2sx2myxgw",
    #     "storyvaloper1wwmwvr4qc8206chufegztl76089wqjjh42vvrl",
    #     "storyvaloper1vcwuam4kp9yq33vcmplkny2rgjcw8tn8ap87pz",
    #     "storyvaloper1gyqunjc9crrwxna3hwvnw6n4laa28r00kxp8pt",
    #     "storyvaloper1re39cmlfjp5cak4x0vhuztf8p0c0rs3r9lzfg3",
    #     "storyvaloper1f2a39llqj4j676htp2xt5tcdj6plh2445elud6",
    #     "storyvaloper1ep7plyu3nx2hrm782zjknjkqzu2hctkfzjnw9q",
    #     "storyvaloper1lslqq5p5x4zs6t8xrt56mkjjg7cazzft787xzg",
    #     "storyvaloper1qxvcwwg608y2gz8q8c94tvtx998wlycg37t9fp",
    #     "storyvaloper1vmvephdv9seepyr9a0zc05208ypu5n0x4k5gwc",
    #     "storyvaloper18r2ufqkmsk0zk5mv3mnfr97hgllyfz4w07pl8y",
    #     "storyvaloper1tjnj658854xdly7fmyugu74rn0z2egtk5u7tp9",
    #     "storyvaloper1k0e9hyrnjkkt702sh94pz048t9pncsd3rawhxn",
    #     "storyvaloper1m8f8kqcjd846pxqn6v3sz78j6t0jczd9uwky4z",
    #     "storyvaloper18waxk5d6zt5k7xnwd5pp8kmsalupthcwdwmr24",
    #     "storyvaloper15xdpesljm34mgllad7wgjfyknxmz8h9yqqzd6y",
    #     "storyvaloper18mwm9q0g0wu926x2xmzc4u4g2dngmyhex26cnx",
    #     "storyvaloper16uszl4wjjjwg9xphn497dpz53k06clthqcajpk",
    #     "storyvaloper1rjn5ew03pr392s0g23rwvnhtnw8z0ctnm0t4gx",
    #     "storyvaloper1v37kgmzm84vau7jul9a8x7yak37u3ckesdvst4",
    #     "storyvaloper1pjex68tfp7yh8uvr674z23ae0sv36zxhqzwzj4",
    #     "storyvaloper16hm543j5ckvzj4u8djkkfftkhtrtklnp0f3427",
    #     "storyvaloper16ln96fdcgfa80yp4a0l8tcyqm9x5lnn4yxuuul",
    #     "storyvaloper18jpaljc3k25wjt96pdaty7nesvt5y30z3afpzl",
    #     "storyvaloper1sgl4s0ysqq8e3fdhrpxnlw7w2xjq6agqvp8kur",
    #     "storyvaloper1z7htzlaktt8drr9nu5uqplxrdxnq345qcsmedy",
    #     "storyvaloper15k0x375pddkjde6h3zxu559wn9v294l726z787",
    #     "storyvaloper1j08gnmvjz5ernv4s4w40jmntwfl7y5hlyxp8jc",
    #     "storyvaloper1a0jfeywlgey88rw9z987phxgr2rrhz3ep2n0rf",
    #     "storyvaloper109xy2m2jsej7p0h64gr6rnl7h9gnjg0j0waufv",
    #     "storyvaloper1jn8fqm7y3kzw0lvc0a8drqlyv6m6vr2rfdhgwu",
    #     "storyvaloper1lfz5gpyss52c4z0tl40c7xzdwdwvksf9xyhmte",
    #     "storyvaloper13xzqljg3yfa5e7ghuvmtzxnwgxyysk9ky3cmzl",
    #     "storyvaloper16ds8xjjpt6dm570zm8c9f4l24lepter3k5u8jy",
    #     "storyvaloper1c564q6073tpthj2x6qhmjrqye4f2txgg9gt2hv",
    #     "storyvaloper1xt07svwfcq70gjyn60ggdk04d8p5mcu9x3emh7",
    #     "storyvaloper1hysw78eqmnfvtgsx7x4eg7gl73raper4qd8q4x",
    #     "storyvaloper13jg060qm37deslld3hjwpzfnwgrxrw75tt0jfh",
    #     "storyvaloper19x42aqxn7ljsd6jm4492gz5c3n6na88vaxmtgj",
    #     "storyvaloper10e4yrqkm6clysrd2dkvkwrtq7u7cphzj7cxay3",
    #     "storyvaloper12956laxe7s5007ye37z7wtdv7h0km5gp4p3xum",
    #     "storyvaloper12sa4vsnhuhmw52sht97uhm3fch57enf2sjydhw",
    #     "storyvaloper1gfky97qzujagey6pnw2xz29tduvc34zrn98gg0",
    #     "storyvaloper1s4cyy5hasyrmnwcx9rerjuwthu38hna9q8vajc"
    # ]

    json_file_path = "metrics-end-day.json"
    output_xlsx_path = "all-metrics.xlsx"
    sheet_name = "Round 1"

    print(len(ROUND_1_VALOPERS))

    # print(len(ROUND_2_VALOPERS))

    with open(json_file_path, "r") as file:
        metrics = json.load(file)

    vals = [val['valoper'] for val in metrics['validators']]


    for val in ROUND_1_VALOPERS:
        if val not in vals:
            print(val)
    cutoff_date = datetime.strptime("2024-12-04", "%Y-%m-%d")

    ROUND_1_DAYS = []
    ROUND_2_DAYS = []

    for date_str in metrics["day_boundaries"]:
        current_date = datetime.strptime(date_str, "%Y-%m-%d")
        
        if current_date <= cutoff_date:
            ROUND_1_DAYS.append(date_str)
            
        else:
            ROUND_2_DAYS.append(date_str)
        
    main(metrics=metrics,
         days=ROUND_1_DAYS,
         output_xlsx_path=output_xlsx_path,
         sheet_name=sheet_name,
         valopers=ROUND_1_VALOPERS
         )




